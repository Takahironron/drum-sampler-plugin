import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

from pathlib import Path
OUT = str(Path(__file__).resolve().parent.parent / "docs" / "research" / "ドラム音源_機能比較.xlsx")
CHECKED = "2026-10-22"

# ---------- sources ----------
S = {
 "S1": ("Toontrack: Superior Drummer 3 製品ページ", "https://www.toontrack.com/product/superior-drummer-3/", "公式"),
 "S2": ("Toontrack ブログ: EZdrummer 3 vs. Superior Drummer 3", "https://www.toontrack.com/blog/ezdrummer-3-vs-superior-drummer-3/", "公式"),
 "S3": ("Toontrack: Superior Drummer 3 リリースノート（3.4.0 2025-04-22、最新 3.4.4 2026-03-26）", "https://www.toontrack.com/release-notes/superior-drummer-3/", "公式"),
 "S4": ("Toontrack: EZdrummer 3 製品ページ", "https://www.toontrack.com/product/ezdrummer-3/", "公式"),
 "S5": ("XLN Audio: Addictive Drums 2 製品ページ", "https://www.xlnaudio.com/products/addictive_drums_2", "公式"),
 "S6": ("XLN サポート: What is Addictive Drums 2?", "https://support.xlnaudio.com/hc/en-us/articles/16382433436701-What-is-Addictive-Drums-2", "公式"),
 "S7": ("XLN サポート: Interface overview - Beats page", "https://support.xlnaudio.com/hc/en-us/articles/16593258822429-Interface-overview-Beats-page", "公式"),
 "S8": ("XLN サポート: Interface overview - The Mixer section", "https://support.xlnaudio.com/hc/en-us/articles/16593352540189-Interface-overview-The-Mixer-section", "公式"),
 "S9": ("XLN サポート: Update: Addictive Drums 2 version 2.6", "https://support.xlnaudio.com/hc/en-us/articles/25294308272413-Update-Addictive-Drums-2-version-2-6-Workflow-improvements-resurrected-features-and-more", "公式"),
 "S10": ("BFD Drums: BFD3 製品ページ", "https://www.bfddrums.com/bfd3/", "公式"),
 "S11": ("BFD3 マニュアル: Using BFD3", "https://internal.bfddrums.com/manuals/BFD30/using_bfd3.htm", "公式マニュアル"),
 "S12": ("BFD3 マニュアル: Grooves page", "https://internal.bfddrums.com/manuals/BFD30/grooves_page.htm", "公式マニュアル"),
 "S13": ("BFD3 マニュアル: Drum editor（Tech／Model）", "https://internal.bfddrums.com/manuals/BFD30/drum_editor.htm", "公式マニュアル"),
 "S14": ("BFD Drums: BFD Player 製品ページ", "https://www.bfddrums.com/bfd-player/", "公式"),
 "S15": ("Steven Slate Drums: SSD 5.5 製品ページ", "https://stevenslatedrums.com/ssd5", "公式"),
 "S16": ("SSD サポート: How to use Groove AI in SSD5.5（2024-11-13）", "https://support.stevenslatedrums.com/hc/en-us/articles/27780053480215-How-to-use-Groove-AI-in-SSD5-5", "公式"),
 "S17": ("SSD サポート: Overview", "https://support.stevenslatedrums.com/hc/en-us/articles/360033795413-Overview", "公式"),
 "S18": ("GetGood Drums: Modern & Massive 製品ページ", "https://ggd.co/products/modern-and-massive", "公式"),
 "S19": ("Native Instruments: Abbey Road Modern Drummer 製品ページ", "https://www.native-instruments.com/en/products/komplete/drums/abbey-road-modern-drummer/", "公式"),
 "S20": ("Native Instruments: Abbey Road Vintage Drummer 製品ページ", "https://www.native-instruments.com/en/products/komplete/drums/abbey-road-vintage-drummer/", "公式"),
 "S21": ("Steinberg: Groove Agent 6 ヘルプ（目次）", "https://www.steinberg.help/r/groove-agent/6.0/en", "公式マニュアル"),
 "S22": ("Steinberg: Groove Agent 6 ヘルプ The Kit（Acoustic Agent）", "https://www.steinberg.help/r/groove-agent/6.0/en/halion/topics/acoustic_agent/the_kit_r.html", "公式マニュアル"),
 "S23": ("Steinberg: Groove Agent SE 6 ヘルプ New Features", "https://www.steinberg.help/r/groove-agent-se/6.0/en/halion/topics/new_features/new_features.html", "公式マニュアル"),
 "S24": ("Steinberg フォーラム: Groove Agent 6 is available now（2026-02-19、公式の告知）", "https://forums.steinberg.net/t/groove-agent-6-is-available-now/1024385", "公式（フォーラム告知）"),
 "S25": ("IK Multimedia: MODO DRUM 1.5 製品ページ", "https://ikmultimedia.com/products/mododrum", "公式"),
 "S26": ("UJAM サポート: Virtual Drummer User Guide", "https://support.ujam.com/hc/en-us/articles/4404979251346-Virtual-Drummer-User-Guide", "公式"),
 "S27": ("Native Instruments: Battery 4 製品ページ", "https://www.native-instruments.com/en/products/komplete/drums/battery-4/", "公式"),
 "S28": ("XLN Audio: XO 製品ページ", "https://www.xlnaudio.com/products/xo", "公式"),
 "S29": ("MusicTech: Algonaut Atlas 2 レビュー（2021-07-28）", "https://musictech.com/reviews/software-instruments/algonaut-atlas-2-review/", "レビュー"),
 "S30": ("Audiomodern: Playbeat 4 製品ページ", "https://audiomodern.com/shop/plugins/playbeat-4/", "公式"),
 "S31": ("Apple サポート: Session Players overview in Logic Pro for Mac", "https://support.apple.com/guide/logicpro/session-players-overview-lgcpbf624405/mac", "公式"),
 "S32": ("Apple サポート: Session Player Editor interface", "https://support.apple.com/guide/logicpro/drummer-editor-interface-lgcpe331a1dc/mac", "公式"),
 "S33": ("Apple サポート: Set Drummer to follow another track's rhythm（Logic Pro 10.7 版の説明）", "https://support.apple.com/guide/logicpro/follow-the-rhythm-of-another-track-lgcp45310225/10.7/mac/11.0", "公式"),
 "S34": ("Apple サポート: Edit a Drummer performance in Logic Pro for Mac", "https://support.apple.com/guide/logicpro/edit-a-drummer-performance-lgcp08ff7ac2/mac", "公式"),
 "S35": ("Apple Newsroom: Logic Pro takes music-making to the next level with new AI features（2024-05-07）", "https://www.apple.com/newsroom/2024/05/logic-pro-takes-music-making-to-the-next-level-with-new-ai-features/", "公式"),
 "S36": ("Magenta: Magenta Studio - Ableton Live Plugin", "https://magenta.tensorflow.org/studio/", "公式"),
 "S37": ("SSD サポート: How To Add 3rd Party Samples In SSD5.5（記事タイトルのみ確認）", "https://support.stevenslatedrums.com/hc/en-us/articles/4403633569559-How-To-Add-3rd-Party-Samples-In-SSD5-5-To-Create-Your-Own-Kit", "公式（タイトルのみ）"),
 "S38": ("BFD3 マニュアル: Using Electronic Drumkits（ページの存在のみ確認）", "https://internal.bfddrums.com/manuals/BFD30/using_electronic_drumkits.htm", "公式マニュアル（タイトルのみ）"),
 "S39": ("Spectrasonics: Stylus RMX Features", "https://www.spectrasonics.net/products/stylusrmx-features.php", "公式"),
 "S40": ("Spectrasonics: Stylus RMX Xpanded Overview（レガシー製品・特別注文の記載）", "https://www.spectrasonics.net/products/stylusrmx/overview.php", "公式"),
 "S41": ("Sugar Bytes: DrumComputer 製品ページ", "https://sugar-bytes.de/drumcomputer", "公式"),
 "S42": ("Serato: Serato Sample 製品ページ", "https://serato.com/sample", "公式"),
 "S43": ("Native Instruments: Maschine 3 製品ページ", "https://www.native-instruments.com/products/maschine-3", "公式"),
 "S44": ("Arturia: Spark 2 製品ページ", "https://www.arturia.com/store/spark2", "公式"),
 "S45": ("Output サポート: Lines, Kits, and Samples（Arcade）", "https://support.output.com/en/articles/10297577-lines-kits-and-samples", "公式"),
 "S46": ("Loopcloud ブログ: Loopcloud 7.4: How to Use Our New MIDI Features（2024-12-11）", "https://www.loopcloud.com/cloud/blog/5251-Loopcloud-7-4-How-to-Use-Our-New-MIDI-Features", "公式"),
 "S47": ("Sonic Charge: Microtonic 製品ページ", "https://soniccharge.com/microtonic", "公式"),
 "S48": ("iZotope サポート: Support Notice for Iris 2, Breaktweaker, and Trash 2（販売終了）", "https://support.izotope.com/hc/en-us/articles/47669652368653-Support-Notice-for-Iris-2-Breaktweaker-and-Trash-2", "公式"),
}

# ---------- features ----------
FEAT = [
 ("F01", "MIDIグルーヴ集（同梱）", "演奏データ（ビート・フィル）のライブラリを持つ"),
 ("F02", "リズム入力でグルーヴ検索", "タップや1小節の打ち込みから、近いグルーヴを探す"),
 ("F03", "曲・録音に合わせた提案／追従", "手持ちのオーディオや MIDI を解析して、合うグルーヴを提案したり追従したりする"),
 ("F04", "パターンの自動生成・変奏", "既存ライブラリの検索ではなく、パターンを生成・変奏する"),
 ("F05", "曲構成づくり", "ソングトラック等で、イントロ〜エンディングの流れを組む"),
 ("F06", "内蔵のMIDI／パターン編集", "グリッドエディタやステップシーケンサーを内蔵"),
 ("F07", "演奏スタイル・複雑さの変形", "既存グルーヴの叩き方・密度・強さを一括で変える"),
 ("F08", "ヒューマナイズ・スウィング・タイトネス", "タイミングやベロシティの揺れ、スウィング量の調整"),
 ("F09", "フィルの制御", "フィルの量・複雑さ・バリエーションを扱う専用の仕組み"),
 ("F10", "ドラム録音→MIDI変換", "ドラムのオーディオから打点を検出して MIDI にする"),
 ("F11", "マルチマイクのミキサー", "クローズ・オーバーヘッド・ルーム等をチャンネル別に混ぜる"),
 ("F12", "ブリード（かぶり）制御", "他の太鼓の音がマイクに入る量を調整する"),
 ("F13", "DAWへの個別アウト", "チャンネルを個別出力で DAW に送る"),
 ("F14", "内蔵エフェクト", "EQ・コンプ・リバーブ等を内蔵"),
 ("F15", "チューニング・ピッチ系", "音程、ベロシティで音程を変える、ピッチエンベロープ等"),
 ("F16", "減衰・エンベロープ・ダンピング", "音の長さや減衰、ミュートの調整"),
 ("F17", "合成・モデル化（サンプル再生以外）", "共鳴・スウェルのモデル化、モーダル合成、シンセ合成など、サンプル再生以外の計算で音や挙動を作る"),
 ("F18", "ハイハット開閉・チョーク", "ハイハットの開き具合、シンバルのチョーク等の制御"),
 ("F19", "ベロシティレイヤー／ラウンドロビン数の明記", "収録の層の数や同じ強さでの差し替え数を公式に記載"),
 ("F20", "電子ドラム対応", "電子ドラム用のマッピングプリセットやベロシティ応答の調整"),
 ("F21", "自前サンプルの読み込み", "ユーザーのサンプルを鳴らせる"),
 ("F22", "キットピースの差し替え・組み合わせ", "別キット・別拡張の太鼓を組み合わせてキットを作る"),
 ("F23", "サンプルの自動整理・分類", "音の類似度や種類で自動的に並べる・分類する"),
 ("F24", "サラウンド／イマーシブ", "サラウンドのマイクで再生・ミックスできる"),
 ("F25", "マクロ・スナップショット", "複数パラメータをまとめる操作、設定の比較（A/B）"),
 ("F26", "MIDI・オーディオの書き出し", "ドラッグ＆ドロップやバウンスで DAW に渡す"),
 ("F27", "AI・機械学習をうたう機能", "開発元が AI／機械学習と説明している機能"),
 ("F28", "オーディオループのテンポ・キー追従", "録音済みのループ（オーディオのグルーヴ）を曲のテンポやキーに合わせて鳴らす"),
 ("F29", "ループのスライス・打ごとの編集", "ループを打ごと・区間ごとに分けて鳴らす、反転するなど"),
 ("F30", "曲からの音源分離（ステム）", "ミックス済みの曲からドラム等を分離する"),
 ("F31", "拍子・フィールの変換", "オーディオのグルーヴを別の拍子やノリに作り替える"),
 ("F32", "音色・キットのランダム生成", "ドラムの音色やキットをランダム・変異で作る"),
]
FIDS = [f[0] for f in FEAT]

# ---------- products ----------
# (id, name, vendor, type, form, price, summary)
PROD = [
 ("P01", "Superior Drummer 3", "Toontrack", "サンプル", "プラグイン／単体", "—（未確認）", "230GB超の収録、最大11.1サラウンド、音から MIDI を作る Tracker、細かいブリード制御"),
 ("P02", "EZdrummer 3", "Toontrack", "サンプル", "プラグイン／単体", "—（未確認）", "オーディオや MIDI を入れるとグルーヴを提案する Bandmate、Tap2Find、Song Creator など作曲向け機能"),
 ("P03", "Addictive Drums 2", "XLN Audio", "サンプル", "プラグイン（VST/AU/AAX）／単体", "—（未確認）", "1小節を描いて探す Grid Search、Beat Transformer、スナップショットの A/B 比較"),
 ("P04", "BFD3（3.5）", "inMusic（BFD Drums）", "サンプル＋モデル化", "プラグイン／単体", "—（未確認）", "タム共鳴・シンバルのスウェルのモデル化、ダンピングなど Tech／Model の細かい操作、ルーディメントで描くグルーヴ編集"),
 ("P05", "BFD Player", "inMusic（BFD Drums）", "サンプル", "プラグイン／単体", "無料", "無料。リードハンド等のタグで探せる340グルーヴ、拡張をまたいだキットピースの差し替え"),
 ("P06", "Steven Slate Drums 5.5", "Steven Slate Drums", "サンプル", "プラグイン", "月US$9.99のレント・トゥ・オウン（一括価格は未確認）", "打ち込みやオーディオから合うグルーヴを探す Groove AI（Step AI／Drop AI）、Morph ノブ"),
 ("P07", "Modern & Massive", "GetGood Drums", "サンプル", "Kontakt（無料 Player 可）", "US$99", "シンバルとスネアのブリードを行列で調整、グルーヴプレイヤー内蔵"),
 ("P08", "Abbey Road Modern Drummer（Vintage Drummer も参照）", "Native Instruments", "サンプル", "Kontakt（無料 Player 可）", "—（未確認）", "最大27ベロシティ層・6ラウンドロビン、スネアのブリードを別録り、Tightness／Swing"),
 ("P09", "Groove Agent 6", "Steinberg", "サンプル＋サンプラー", "プラグイン／単体", "—（未確認）", "2026-02-19 発売の告知 [S24]。Acoustic／Beat／Percussion の3種のエージェント、音の自動分類、Decompose（音程成分と雑音成分に分解）"),
 ("P10", "MODO DRUM 1.5", "IK Multimedia", "モーダル合成＋サンプル", "プラグイン（VST2/VST3/AU/AAX）／単体", "US$199.99（SE US$149.99、CS 無料）", "寸法・材質・ヘッド・スティック・叩き方をモデルで変える"),
 ("P11", "Virtual Drummer（各タイトル）", "UJAM", "サンプル＋フレーズ再生", "プラグイン", "—（未確認）", "鍵盤でフレーズを切り替え、強さを5段階とモジュレーションホイールで変える。ミックスは Amount で効き具合を調整"),
 ("P12", "Battery 4", "Native Instruments", "サンプラー", "プラグイン／単体", "—（未確認）", "セル行列のドラムサンプラー。波形編集、バス経路のドラッグ＆ドロップ、サイドチェイン"),
 ("P13", "XO", "XLN Audio", "サンプラー＋シーケンサー", "プラグイン／単体", "—（未確認）", "手持ちのワンショットを音の類似度で地図状に並べる。Accentuator、グルーヴテンプレート"),
 ("P14", "Atlas 2", "Algonaut", "サンプラー＋シーケンサー", "プラグイン／単体", "—（未確認）", "サンプルを解析して種類分けし、似た音を2次元の地図にまとめる。キットの自動作成"),
 ("P15", "Playbeat 4", "Audiomodern", "パターン生成＋サンプル", "プラグイン（Logic 用 MIDI FX 版、iPad AUv3 あり）", "€89（旧版から €29）", "ジャンルに基づく AI のパターン生成、リミックスの生成と保存"),
 ("P16", "Logic Pro の Drummer（Session Players）", "Apple", "DAW 内蔵（演奏生成）", "Logic Pro の機能", "Logic Pro に付属", "別トラックのリズムに追従（Follow）、複雑さと強さ、フィル量・複雑さ、再生成"),
 ("P17", "Magenta Studio 2", "Google（Magenta）", "機械学習による MIDI 生成", "Ableton Live 用デバイス（Live 10.1 Suite 以上）", "—（ページに記載なし）", "ベースやメロディからドラムを作る Drumify、実演奏で学習した Groove（揺れ付け）、2つのビートの補間"),
 ("P18", "Stylus RMX Xpanded", "Spectrasonics", "オーディオループ（グルーヴ）", "プラグイン", "特別注文品（レガシー製品、サポートは継続）。価格は未確認", "オーディオのグルーヴをテンポ同期で鳴らし、Chaos Designer で変奏、Time Designer で拍子・ノリを変える。REX で自前ループを取り込む"),
 ("P19", "DrumComputer", "Sugar Bytes", "合成（3層）＋サンプラー", "プラグイン／単体（iPad 版あり）", "—（ページの US$129 が iPad 版かデスクトップ版か判別できず）", "共鳴体・オシレーター・再合成/サンプラーの3層の音源、確率やロールのシーケンサー、Remix と Auto Fill、音色のランダム生成"),
 ("P20", "Serato Sample", "Serato", "サンプラー（チョップ）", "プラグイン", "月額の Producer Suite か買い切り（金額は未確認）", "曲から実時間でドラム等のステムを分離し、Find Samples でキューポイントを自動配置してパッドで鳴らす"),
 ("P21", "Maschine 3（ソフトウェア）", "Native Instruments", "ビート制作環境（サンプラー＋ドラムシンセ）", "ソフトウェア（専用ハードと連携）", "US$109（Maschine Central 込み）", "曲をベース・ドラム・ボーカル等に分離、サンプリング、ドラムシンセ、シーンごとのテンポ"),
 ("P22", "Spark 2", "Arturia", "アナログモデリング合成・物理モデル・サンプル", "プラグイン／単体", "—（未確認）", "TAE のアナログ合成と物理モデルとサンプルの音源、64ステップのシーケンサー、ループモード（Divide／Move）、モジュラー式のドラムシンセ編集"),
 ("P23", "Arcade", "Output", "オーディオループ（サブスクリプションの素材）", "プラグイン", "サブスクリプション（金額は未確認）", "15本のループを白鍵に並べ、曲のテンポとキーに合わせて鳴らす。Line ごとにマクロと Modifier を用意"),
 ("P24", "Loopcloud", "Loopmasters", "素材の検索・再生（ループ・MIDI）", "プラグイン／単体", "—（未確認）", "15万以上の MIDI ファイルを検索し、テンポとキーに合わせて試聴。パートごとの書き出し"),
 ("P25", "Microtonic", "Sonic Charge", "シンセ合成（サンプル不使用）", "プラグイン", "US$99", "サンプルを使わない合成のドラム、ホスト同期のパターンエンジン、モーフィングとランダム化、コンピューター生成のパターン集 Patternarium"),
]
PIDS = [p[0] for p in PROD]

# ---------- cells: (product, feature) -> (mark, detail, sources) ----------
# mark: ○ = 確認できた, △ = 近い機能はあるが条件付き／部分的
C = {}
def c(p, f, mark, text, src):
    C[(p, f)] = (mark, text, src)

# P01 Superior Drummer 3
c("P01","F01","○","約1,640 の MIDI ファイル。曲のパート別に分類","S1")
c("P01","F02","○","Tap2Find（EZdrummer 3 と共通）","S2")
c("P01","F05","○","Song Track（テンポ・拍子変更、複数バージョン）、Song Creator","S1,S2")
c("P01","F06","○","MIDI グリッドエディタ。nudge・random などの高度な編集","S1")
c("P01","F07","○","Edit Play Style","S1")
c("P01","F08","○","スウィング・クオンタイズ、random。電子ドラム入力のヒューマナイズ","S1")
c("P01","F10","○","Tracker: ドラムのオーディオを解析して MIDI に変換","S1")
c("P01","F11","○","22チャンネルのミキサーと32の内部バス。全マイクが個別チャンネル","S1,S2")
c("P01","F12","○","太鼓ごとのブリード制御（Individual bleed controls、Instrument bleed）","S1,S2")
c("P01","F14","○","低負荷のインサートエフェクト35種","S1")
c("P01","F15","○","ヘッドとシンバルのチューニング、ピッチエンベロープ","S1")
c("P01","F16","△","リバース再生の記載はあるが、減衰の操作は未確認","S1")
c("P01","F18","○","ハイハットペダルの応答と開き具合の編集（電子ドラム設定）","S1")
c("P01","F20","○","主要メーカーのマッピングプリセット、Snare Zone Control、MIDI 出力","S1")
c("P01","F21","○","自前サンプルの読み込み","S1")
c("P01","F22","○","ドラムスタック作成、SDX・EZX ライブラリの太鼓を混在","S1")
c("P01","F24","○","最大11.1サラウンド（追加のルームマイク11本）","S1")
c("P01","F25","○","マクロコントロール（3.4.0 でプラグインのパラメータ名をマクロ名に合わせた）","S1,S3")
c("P01","F26","○","16／24bit のオーディオバウンス","S1")

# P02 EZdrummer 3
c("P02","F01","○","2,500 以上のグルーヴとフィル。曲構成のパート別。導入・終止用の素材も","S4")
c("P02","F02","○","Tap2Find: リズムを叩くと近いグルーヴを一覧","S4")
c("P02","F03","○","Bandmate: オーディオか MIDI を入れるとグルーヴを提案（Bandmate は EZdrummer 3 のみと公式ブログ）","S4,S2")
c("P02","F05","○","Song Creator、Song Track でパートを並べ替え","S4")
c("P02","F06","○","グリッドエディタ（Superior Drummer 3 と同じ）","S4,S2")
c("P02","F07","○","Edit Play Style（演奏の実時間変更）","S4")
c("P02","F11","○","ミキサー。Michael Ilbert のマイク設定のプリセット","S4")
c("P02","F14","○","プリセットごとのエフェクト操作","S2")
c("P02","F18","○","柔軟なハイハット制御（電子ドラム）","S4")
c("P02","F20","○","電子ドラムのプリセット、ベロシティ応答の調整、直接録音","S4")

# P03 Addictive Drums 2
c("P03","F01","○","Beats & Fills の MIDI ライブラリ。スタイル・テンポ・拍子で絞り込み","S6,S7")
c("P03","F02","○","Grid Search: 1小節16分のパターンを描いて探す。Similar で選択中のビートから検索","S7")
c("P03","F07","○","Beat Transformer: Range、Accents、Random、キットピースの差し替え、Speed・Length","S7")
c("P03","F08","○","Random ノブでタイミングとベロシティを別々に揺らす","S7")
c("P03","F09","○","Fills の絞り込み、Variant List でバリエーションを選ぶ","S7")
c("P03","F11","○","クローズマイクのモノラル10ch＋OH・ルームのステレオ2ch","S8")
c("P03","F13","○","全チャンネルを個別出力に。2.6 で Route All Channels","S8,S9")
c("P03","F14","○","チャンネルごとに EQ・コンプ・歪み・トランジェント・テープ。センドの Delerb 2系統","S6,S8")
c("P03","F15","○","Pitch","S6")
c("P03","F16","○","音量エンベロープ","S6")
c("P03","F20","○","電子ドラムは設定不要と公式。電子ドラムからの MIDI 録音","S5,S7")
c("P03","F21","△","Trigger キットピース（追加のワンショットを重ねる）。任意のサンプル読み込みかは未確認","S6")
c("P03","F22","○","ライブラリ内のキットピースを組み合わせてキットを作る","S5")
c("P03","F25","○","Snapshots: 設定を最大5個保存して A/B 比較（2.6）","S9")
c("P03","F26","○","Audio Recorder でループを作り DAW へドラッグ（2.6）","S9")

# P04 BFD3
c("P04","F01","○","82 の Groove Palette","S10")
c("P04","F05","○","Drum Track にグルーヴを並べて曲全体を作る","S12")
c("P04","F06","○","Groove Editor。ルーディメントから描く","S10,S12")
c("P04","F07","○","Groove FX で dynamics と complexity も調整（詳細の記載なし）","S12")
c("P04","F08","○","Groove FX の Humanize・Swing（非破壊・実時間）。Model ページの Variance","S12,S13")
c("P04","F11","○","ダイレクトマイク＋OH・Room・Amb3。Microphone Inspector","S10,S11")
c("P04","F12","○","Bleed の Trim・Send・Return、行き先（キック・スネア）を指定","S13")
c("P04","F13","○","チャンネルを個別出力へ。外部信号でコンプのサイドチェイン","S10")
c("P04","F14","○","48種（AIR の13種を含む）","S10")
c("P04","F15","○","Tune（最大1オクターブ）、Vel to Pitch、タムを鍵盤で合わせる","S13")
c("P04","F16","○","Damping（量・周波数・アンビエントへの比率）、Vel to Damp、エンベロープ","S13")
c("P04","F17","○","タム共鳴とブリードのモデル化、シンバルのスウェル、ダイナミクスの整合","S10,S13")
c("P04","F18","○","Hihat tighten（閉じるほど減衰・音程上昇）、チョークの応答とチョークグループ","S13")
c("P04","F20","△","マニュアルに電子ドラム用のページがある（内容は未確認）","S38")
c("P04","F26","○","マルチチャンネルのオーディオ書き出し、グルーヴのドラッグ書き出し","S10")

# P05 BFD Player
c("P05","F01","○","340 グルーヴ。ジャンル・テンポ・リードハンドのタグで検索","S14")
c("P05","F11","○","プリセット付きのミキサー","S14")
c("P05","F13","○","ドラム出力を個別にルーティング","S14")
c("P05","F22","○","別の拡張パックのキットピースに差し替え","S14")
c("P05","F26","○","グルーヴをタイムラインへドラッグ","S14")

# P06 SSD5.5
c("P06","F01","○","2,400 以上の MIDI グルーヴ","S15")
c("P06","F02","○","Groove AI の Step AI: 打ち込んだパターンに合うグルーヴを探す","S16")
c("P06","F03","○","Groove AI の Drop AI: 読み込んだオーディオに合うグルーヴを探す。曲のトランジェントを解析","S16,S15")
c("P06","F07","○","Morph ノブでグルーヴを調整（製品ページの FAQ）","S15")
c("P06","F11","○","ミキサーでルームとクローズマイクを調整","S17")
c("P06","F13","○","マルチアウトで DAW に個別に送る","S17")
c("P06","F14","○","Mix Console の Effects Rack","S15")
c("P06","F15","○","チューニング、アーティキュレーション、ベロシティ","S15")
c("P06","F16","○","ADSR","S15")
c("P06","F21","○","サードパーティのサンプルを加えてキットを作る（サポート記事の題名）","S37")
c("P06","F22","○","太鼓を選んでキットを作る。最大32インストゥルメント","S17")
c("P06","F27","○","Groove AI","S16")

# P07 GGD Modern & Massive
c("P07","F01","△","グルーヴプレイヤー内蔵（収録数は未確認）","S18")
c("P07","F12","○","シンバルとスネアのブリードの行列","S18")
c("P07","F14","○","内蔵の処理とリバーブ","S18")

# P08 Abbey Road Modern Drummer
c("P08","F01","○","1,000 グルーヴ・5ジャンル（Vintage は 600・8ジャンル）","S19,S20")
c("P08","F08","○","Tightness（グリッド通り〜ゆるさ）と Swing","S19,S20")
c("P08","F09","○","フィルとバリエーションを収録","S20")
c("P08","F11","○","スネア上下、OH、ステレオルーム、特殊マイクのフェーダー","S19")
c("P08","F12","○","スネアのブリードを別録りし量を調整","S19,S20")
c("P08","F14","○","4バンド EQ、コンプ、Transient Master、テープ、コンボリューションリバーブ（IR 30種）","S19")
c("P08","F19","○","最大27ベロシティ層、同じ強さで最大6ラウンドロビン（Vintage は4）。左右の手を別収録","S19,S20")
c("P08","F20","○","GM、V-Drums、Superior Drummer、BFD、Addictive Drums 等のマッピングテンプレート","S19")
c("P08","F26","○","グルーヴを MIDI でホストにドラッグ","S19")

# P09 Groove Agent 6
c("P09","F01","△","Style Player・MIDI Player を搭載（同梱スタイルの数は未確認）","S21")
c("P09","F06","○","Pattern Player: ステップシーケンサー。スタイルや MIDI をパターンに変換して編集","S23")
c("P09","F11","○","太鼓ごとの Room・Overhead レベル。全エージェント共通のミキサー","S22,S23")
c("P09","F14","○","全エージェントで全エフェクトを使える。6 で EQ・歪み・コンプを追加","S23")
c("P09","F15","○","Tune。Beat Agent の Pitch Detection と Pitch Quantize","S22,S23")
c("P09","F16","○","Attack・Hold・Decay","S22")
c("P09","F18","○","シンバルのチョーク（ノートオフ・アフタータッチ）、ハイハットの開きの最小・最大、CC の遮断","S22")
c("P09","F21","○","Sample タブ、単体版でのサンプル録音、ループのスライス、Decompose（音程成分と雑音成分に分解）","S21,S23")
c("P09","F23","○","ドラム音の自動分類（Automatically Classifying Drum Sounds）","S21")

# P10 MODO DRUM
c("P10","F01","○","1,400 以上の MIDI グルーヴ","S25")
c("P10","F11","○","スタジオ品質の処理付きミキサー","S25")
c("P10","F14","○","19種のプロセッサー、トラック・バス・マスターに8バンド EQ","S25")
c("P10","F15","○","チューニング","S25")
c("P10","F17","○","実時間のモーダル合成＋サンプル。寸法・材質・ヘッド・スティック・叩き方","S25")
c("P10","F26","○","グルーヴを DAW へドラッグ＆ドロップ","S25")

# P11 UJAM Virtual Drummer
c("P11","F01","○","30スタイル。イントロ・Verse・Chorus・エンディング等のパート別","S26")
c("P11","F07","○","Verse・Chorus に強さ順の5バリエーション。モジュレーションホイールで強さを連続変化","S26")
c("P11","F08","○","Humanize（収録の揺れ〜グリッド）、Swing","S26")
c("P11","F09","○","3種のフィル。ループに戻るときにクラッシュを自動追加","S26")
c("P11","F11","△","楽器・グループ別のチャンネルとミックスプリセット（Amount で効き具合）。マイク別ではない","S26")
c("P11","F15","○","チャンネルごとのチューニング","S26")
c("P11","F16","○","チャンネルごとの Decay","S26")
c("P11","F26","○","パートを MIDI として DAW にドラッグ","S26")

# P12 Battery 4
c("P12","F13","△","バスへのドラッグ＆ドロップの経路設定、4バスとマスターでサイドチェイン（DAW への個別出力は未確認）","S27")
c("P12","F14","○","SOLID EQ、SOLID BUS COMP、TRANSIENT MASTER、テープ、LoFi、コンボリューションリバーブ","S27")
c("P12","F16","○","波形上で開始・終了点、音量カーブ、ループ、エンベロープ","S27")
c("P12","F21","○","自前サンプルを読み込みセル行列でキットを作る","S27")

# P13 XO
c("P13","F01","○","240 以上の編集できるビートプリセット","S28")
c("P13","F08","○","グルーヴテンプレート、パターンの nudge、Accentuator","S28")
c("P13","F21","○","自前サンプルの読み込み（重複や散らばったフォルダも整理）","S28")
c("P13","F23","○","ワンショットを音の類似度で視覚的に整理","S28")
c("P13","F26","○","単体サンプル・キット・ステム・ビートをオーディオか MIDI でドラッグ書き出し","S28")
c("P13","F27","△","ページの説明文に「AI-powered drum sampling」とあるのみ。仕組みの説明はない","S28")

# P14 Atlas 2
c("P14","F04","○","キットの自動作成（押すたびに別の組み合わせ）、パターンの自動バリエーション","S29")
c("P14","F06","○","シーケンサー（レーンごとの長さ、回転、nudge、ミラー）","S29")
c("P14","F21","○","フォルダやディスクを解析して読み込む","S29")
c("P14","F23","○","キック・スネア等に分類し、似た音を2次元の地図にまとめる","S29")
c("P14","F26","○","パターンを MIDI かオーディオで書き出し","S29")
c("P14","F27","△","サンプル解析・分類（レビューによる。2021 年の版）","S29")

# P15 Playbeat 4
c("P15","F04","○","ジャンルに基づくパターン生成。スタイルと複雑さを選ぶ。リミックスを無制限に生成・保存","S30")
c("P15","F21","○","自前の音・プリセット・タグを追加（体験版では無効）","S30")
c("P15","F26","△","MIDI 出力部で他のプラグインや外部機器を鳴らす（MIDI ファイル書き出しの記載はない）","S30")
c("P15","F27","○","AI によるパターン生成とうたう","S30")

# P16 Logic Drummer
c("P16","F03","○","Follow: キックとスネアを別トラックのリズムに追従（パーカッションのジャンルは不可）","S33,S31")
c("P16","F04","○","現在の設定で演奏を再生成","S31")
c("P16","F07","○","Complexity と Intensity のスライダー。キット別のパターン選択","S32,S34")
c("P16","F08","○","Swing ノブ","S32")
c("P16","F09","○","Fills Amount と Fill Complexity","S32")
c("P16","F18","○","Hi-Hat の開閉の度合い（Details）","S34")
c("P16","F22","○","Drum Kit Designer でキットピースを交換","S31")
c("P16","F27","○","Session Players を「AI-driven backing band」と説明（2024-05）","S35")

# P17 Magenta Studio
c("P17","F03","○","Drumify: ベース・メロディ・タップしたリズムからドラムを作る","S36")
c("P17","F04","○","Generate（4小節を新規生成）、Continue（最大32小節延長）、Interpolate（2つのビートを補間）","S36")
c("P17","F08","○","Groove: 実際のドラマーの録音で学習したモデルでタイミングとベロシティを付ける","S36")
c("P17","F27","○","Magenta のオープンソースのモデルを利用","S36")

# P18 Stylus RMX
c("P18","F01","○","14GB のコアライブラリ。付属の Xpander 5種と合わせて約10,000 のグルーヴと音","S39")
c("P18","F04","○","Chaos Designer: オーディオのグルーヴに変奏や「即興」を加える。Capture で MIDI 化","S39")
c("P18","F13","○","8ch ミキサー（ミュート・ソロ、Aux センド4系統、出力の割り当て）","S39")
c("P18","F14","○","統合エフェクトラック（27種）","S39")
c("P18","F21","○","ReCycle の REX ファイルで自前のループを取り込む","S39,S40")
c("P18","F26","○","MIDI ファイルをホストにドラッグ","S39")
c("P18","F28","○","グルーヴを実時間でビートに自動同期","S39,S40")
c("P18","F29","○","Edit Groups でループ内の個々の打を非破壊で操作（特定スライスの反転など）","S39,S40")
c("P18","F31","○","Time Designer: オーディオのグルーヴを任意の拍子・フィール・パターンに実時間で変える","S39")

# P19 DrumComputer
c("P19","F06","○","16ステップ×16パターン。トラックごとのテンポ・方向・長さ、確率・ロール・ステップ遅延のモジュレーション","S41")
c("P19","F09","○","Auto Fill（曲の位置でフィルとパラメータ変化）、Remix コントロールでフィル","S41")
c("P19","F14","○","フィルター、コンプ、オーバードライブ、リバーブ2種、Finalizer","S41")
c("P19","F17","○","8つの音源。共鳴体、ウェーブテーブル／アナログのオシレーター、再合成／サンプラーの3層","S41")
c("P19","F21","○","Sampler Mode で自前の音を鳴らす","S41")
c("P19","F26","○","トリガー・モジュレーション・Auto Fill・CC を含む MIDI ファイルをドラッグで書き出し","S41")
c("P19","F27","○","Claude 用の拡張: 文章の指示1つから DrumComputer の音色を生成（製品版が必要）","S41")
c("P19","F32","○","プロファイル付きのランダマイザー、MakeKit でキットを生成","S41")

# P20 Serato Sample
c("P20","F15","○","テンポを変えずにキーを半音・微分音で変更","S42")
c("P20","F21","○","任意の曲やサンプルを読み込んでパッドに割り当て","S42")
c("P20","F23","△","Find Samples: 波形を解析して良い位置にキューポイントを自動配置（分類ではない）","S42")
c("P20","F28","○","パッドごとに音程を保ったまま再生速度を変える","S42")
c("P20","F29","○","最大32のキューポイントを32パッドに割り当て","S42")
c("P20","F30","○","アカペラ・メロディ・ベース・ドラムのステムを実時間で分離","S42")

# P21 Maschine 3
c("P21","F06","○","パターンの録音・編集、ノートの nudge、モジュレーションエディタ","S43")
c("P21","F17","○","内蔵のドラムシンセ","S43")
c("P21","F21","○","サンプルの録音と編集（Sampler ビュー）","S43")
c("P21","F30","○","曲を読み込みベース・ドラム・ボーカル・その他に分離（Sampler ビューでも）","S43")

# P22 Spark 2
c("P22","F06","○","64ステップのシーケンサー","S44")
c("P22","F08","○","Shuffle と Shift でグルーヴをヒューマナイズ","S44")
c("P22","F15","○","パッドとステップボタンで弾ける16音程モード","S44")
c("P22","F17","○","TAE のアナログ合成と物理モデル（多層サンプルと併用）。モジュラー式のドラムシンセ編集","S44")
c("P22","F18","○","チョークグループ","S44")
c("P22","F25","○","オシレーターとフィルター用のカスタムマクロのモジュール","S44")

# P23 Arcade
c("P23","F21","○","自前のサンプルをドラッグで追加","S45")
c("P23","F25","△","Kit ごとにマクロと Modifier を用意（中身の詳細は未確認）","S45")
c("P23","F28","○","ループを曲のテンポとキーに合わせて再生。15本を白鍵 C2〜C4 に配置","S45")

# P24 Loopcloud
c("P24","F01","△","15万以上の MIDI ファイルを検索（ドラム専用ではない）。「Drum kit」の音で試聴","S46")
c("P24","F26","○","Processed Separate Files でパートを個別に書き出し","S46")
c("P24","F28","○","テンポに自動で合わせる、Pitch Lock でキーを固定して試聴","S46")

# P25 Microtonic
c("P25","F04","○","Patternarium: コンピューターが生成した数千のプリセット（パターン）をオンラインで試聴・取得","S47")
c("P25","F06","○","ホストに同期するパターン式のドラムマシン","S47")
c("P25","F17","○","サンプルや事前に作った波形を使わない合成","S47")
c("P25","F32","○","モーフィング、smart alteration、ランダム化","S47")

for (p, f) in C: assert p in PIDS and f in FIDS, (p, f)

# ---------- styles ----------
FONT = "Arial"
hfill = PatternFill("solid", fgColor="305496")
hfont = Font(name=FONT, bold=True, color="FFFFFF")
nfont = Font(name=FONT)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
fill_o = PatternFill("solid", fgColor="C6EFCE")
fill_t = PatternFill("solid", fgColor="FFEB9C")

def header(ws, row, vals):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=row, column=i, value=v)
        x.font = hfont; x.fill = hfill; x.alignment = center; x.border = border

def body(ws, row, vals, align=wrap):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=row, column=i, value=v)
        x.font = nfont; x.alignment = align; x.border = border

wb = openpyxl.Workbook()

# 注記
ws = wb.active; ws.title = "注記"
header(ws, 1, ["項目", "内容"])
notes = [
 ("目的", "ドラム音源プラグイン（サンプルベース中心、モデリング、AI・パターン生成系、ループ／グルーヴ系・ビート制作系を含む。P18〜P25 は 2026-10-22 に追加）の「特徴的な機能」を製品ごと・機能ごとに整理する。"),
 ("確認日", CHECKED + "。各製品の公式ページ・公式マニュアル・公式サポート記事を中心に確認。公式で足りないものはレビュー記事を使い、出典の種別に明記した。"),
 ("記号", "○＝確認した出典に記載がある ／ △＝近い機能はあるが部分的・条件付き、または中身が未確認 ／ 空欄＝確認した出典に記載が見つからない（機能が無いという意味ではない）"),
 ("推測しないこと", "空欄は「無い」ではなく「確認した出典に書かれていない」。例: Superior Drummer 3 のラウンドロビン数、Addictive Drums 2 のブリード操作は、確認したページに記載がなく空欄にしている。"),
 ("価格", "確認したページに価格があったものだけ記載。それ以外は「未確認」。"),
 ("シート", "製品一覧／製品×機能（マトリクス。セルのコメントに詳細）／機能別一覧（機能ごとに持つ製品と中身）／詳細（製品×機能の全行）／出典／ジャンル別キット（lo-fi・psybient）"),
 ("対象外にしたもの", "Perfect Drums（Naughty Seal Audio）、MT Power DrumKit 2、MixWave: 公式の機能説明を取得できず、販売店・KVR・ブログの情報しかなかったため。iZotope BreakTweaker: 公式サイトで販売終了と告知されているため（S48）。Groove Agent 6 の公式「新機能」ページは JavaScript で描画され本文を取得できなかったため、ヘルプで確認した。"),
 ("前回の比較シートとの関係", "物理モデリング中心の比較は「ドラム音源プラグイン比較.xlsx」（D1〜D3）。MODO DRUM と SSD5.5 は両方に入っている。"),
]
for r, (a, b) in enumerate(notes, 2):
    body(ws, r, [a, b])
ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 120

# 製品一覧
ws = wb.create_sheet("製品一覧")
header(ws, 1, ["#", "製品", "開発元", "方式", "形態", "価格（" + CHECKED + "確認）", "特徴的な機能（要約）", "確認できた機能の数", "出典"])
for r, p in enumerate(PROD, 2):
    srcs = sorted({s for (pp, f), v in C.items() if pp == p[0] for s in v[2].split(",")} | ({"S24"} if p[0]=="P09" else set()), key=lambda s: int(s[1:]))
    body(ws, r, list(p) + [None, " ".join(srcs)])
    ws.cell(row=r, column=8, value=f"=COUNTIF('製品×機能'!C{r+1}:{L(2+len(FEAT))}{r+1},\"○\")").alignment = center
widths = [6, 30, 22, 18, 26, 24, 60, 10, 28]
for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
ws.freeze_panes = "C2"

# 製品×機能
ws = wb.create_sheet("製品×機能")
ws.cell(row=1, column=1, value="凡例: ○＝出典に記載あり、△＝部分的・条件付き、空欄＝確認した出典に記載なし。セルのコメントに中身と出典番号。").font = Font(name=FONT, italic=True)
header(ws, 2, ["#", "製品"] + [f"{fid}\n{name}" for fid, name, _ in FEAT])
for r, p in enumerate(PROD, 3):
    body(ws, r, [p[0], p[1]])
    for j, fid in enumerate(FIDS, 3):
        x = ws.cell(row=r, column=j); x.border = border; x.alignment = center; x.font = nfont
        if (p[0], fid) in C:
            mark, text, src = C[(p[0], fid)]
            x.value = mark
            x.fill = fill_o if mark == "○" else fill_t
            cm = Comment(f"{text}\n出典: {src}", "調査"); cm.width = 300; cm.height = 120
            x.comment = cm
last = 2 + len(PROD)
rr = last + 1
ws.cell(row=rr, column=2, value="○ の数").font = Font(name=FONT, bold=True)
ws.cell(row=rr + 1, column=2, value="△ の数").font = Font(name=FONT, bold=True)
for j in range(3, 3 + len(FEAT)):
    ws.cell(row=rr, column=j, value=f'=COUNTIF({L(j)}3:{L(j)}{last},"○")').alignment = center
    ws.cell(row=rr + 1, column=j, value=f'=COUNTIF({L(j)}3:{L(j)}{last},"△")').alignment = center
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 30
for j in range(3, 3 + len(FEAT)): ws.column_dimensions[L(j)].width = 11
ws.row_dimensions[2].height = 90
ws.freeze_panes = "C3"

# 機能別一覧
ws = wb.create_sheet("機能別一覧")
header(ws, 1, ["#", "機能", "説明", "○ の製品数", "△ の製品数", "持つ製品と中身（○）", "部分的（△）"])
pname = {p[0]: p[1] for p in PROD}
for r, (fid, name, desc) in enumerate(FEAT, 2):
    col = L(3 + FIDS.index(fid))
    o = [f"・{pname[p]}: {C[(p, fid)][1]} [{C[(p, fid)][2]}]" for p in PIDS if (p, fid) in C and C[(p, fid)][0] == "○"]
    t = [f"・{pname[p]}: {C[(p, fid)][1]} [{C[(p, fid)][2]}]" for p in PIDS if (p, fid) in C and C[(p, fid)][0] == "△"]
    body(ws, r, [fid, name, desc, None, None, "\n".join(o) or "（確認できた製品なし）", "\n".join(t)])
    ws.cell(row=r, column=4, value=f"='製品×機能'!{col}{rr}").alignment = center
    ws.cell(row=r, column=5, value=f"='製品×機能'!{col}{rr+1}").alignment = center
for i, w in enumerate([6, 26, 36, 9, 9, 90, 60], 1): ws.column_dimensions[L(i)].width = w
ws.freeze_panes = "C2"

# 詳細
ws = wb.create_sheet("詳細")
header(ws, 1, ["製品#", "製品", "機能#", "機能", "記号", "中身", "出典"])
fname = {f[0]: f[1] for f in FEAT}
r = 2
for p in PIDS:
    for fid in FIDS:
        if (p, fid) in C:
            m, t, s = C[(p, fid)]
            body(ws, r, [p, pname[p], fid, fname[fid], m, t, s]); r += 1
for i, w in enumerate([7, 30, 7, 28, 6, 80, 14], 1): ws.column_dimensions[L(i)].width = w
ws.auto_filter.ref = f"A1:G{r-1}"
ws.freeze_panes = "A2"

# 出典
ws = wb.create_sheet("出典")
header(ws, 1, ["番号", "題名", "URL", "種別", "確認日"])
for r, (k, (t, u, kind)) in enumerate(S.items(), 2):
    body(ws, r, [k, t, u, kind, CHECKED])
    ws.cell(row=r, column=3).hyperlink = u
    ws.cell(row=r, column=3).font = Font(name=FONT, color="0563C1", underline="single")
for i, w in enumerate([7, 70, 80, 18, 12], 1): ws.column_dimensions[L(i)].width = w

# ジャンル別（lo-fi・psybient）
GS = {
 "G1": ("XLN Audio: XO Expansion Lo-Fi Beats", "https://www.xlnaudio.com/products/xo/expansion/lo-fi_beats"),
 "G2": ("Native Instruments: Faded Reels", "https://www.native-instruments.com/en/products/komplete/expansions/faded-reels/"),
 "G3": ("Native Instruments: Midnight Flow", "https://www.native-instruments.com/en/products/komplete/expansions/midnight-flow/"),
 "G4": ("Native Instruments: Drift Theory", "https://www.native-instruments.com/en/products/komplete/expansions/drift-theory/"),
 "G5": ("Native Instruments: Opaline Drift", "https://www.native-instruments.com/products/opaline-drift?variant=54051429482799"),
 "G6": ("Toontrack: Lo-Fi EZmix Pack", "https://www.toontrack.com/product/lo-fi-ezmix-pack/"),
 "G7": ("Loopmasters: Psytronica（Famous Audio）", "https://www.loopmasters.com/genres/138-Electronica/products/7576-Psytronica"),
 "G8": ("Plugin Boutique: BigKick Expansion - PsiFi Kicks", "https://www.pluginboutique.com/product/1-Instruments/70-Expansion-Packs/4466-BigKick-Expansion-PsiFi-Kicks"),
 "G9": ("KVR Marketplace: Neopsy Kick（Eplex7 DSP）", "https://www.kvraudio.com/marketplace/neopsy-kick-by-eplex7-dsp"),
 "G10": ("Native Instruments: Electronica & Experimental のエクスパンション一覧", "https://native-instruments.com/en/catalog/maschine/expansions/electronica-experimental"),
}
GENRE = [
 ("lo-fi", "Lo-Fi Beats", "XLN Audio", "XO（P13）の拡張", "編集できるビートプリセット75、ジャンル専用サンプル150以上。chillhop 向け", "US$39（XO 本体 US$149 が別途必要）", "○ 明記", "G1"),
 ("lo-fi", "Faded Reels", "Native Instruments", "Maschine（P21）・Battery 4（P12）のエクスパンション", "Maschine キット50＋Battery キット59。「lo-fi hip hop」と明記", "US$49", "○ 明記", "G2"),
 ("lo-fi", "Midnight Flow", "Native Instruments", "Maschine・Battery 4 のエクスパンション", "Maschine キット50＋Battery キット58。「Chill lo-fi hip hop kits」", "US$49", "○ 明記", "G3"),
 ("lo-fi", "Drift Theory", "Native Instruments", "Maschine・Battery 4 のエクスパンション", "Maschine キット50＋Battery キット59。「Lo-fi kits and patterns」", "US$49", "○ 明記", "G4"),
 ("lo-fi（参考）", "Lo-Fi EZmix Pack", "Toontrack", "EZmix 2／3 のプリセット（ドラム音源ではない）", "lo-fi 用のエフェクトチェーン50、うちドラムキット用11。EZdrummer 等の音を lo-fi に加工する用途", "US$39", "△ 加工用", "G6"),
 ("downtempo（参考）", "Opaline Drift", "Native Instruments", "Maschine・Battery 4 のエクスパンション", "Battery キット50＋Maschine キット45。downtempo electronic 向け（psybient の明記はない）", "US$49", "△ 近いジャンル", "G5"),
 ("psybient", "Psytronica（Famous Audio）", "Loopmasters", "WAV のサンプルパック（Loopcloud（P24）からも単体で入手可）", "ドラムループ61、ワンショット26。対象ジャンルに「Psybient, Psychill」と明記。プラグインのキットではない", "US$46.95", "△ サンプルパック", "G7"),
 ("psytrance（参考）", "BigKick Expansion - PsiFi Kicks", "Plugin Boutique", "BigKick（キック生成プラグイン）の拡張", "キックのプリセット64、アタック32。Trance・Psytrance・Goa 向け", "US$15（確認時はセールで US$5）", "△ キックのみ・別ジャンル", "G8"),
 ("psytrance（参考）", "Neopsy Kick", "Eplex7 DSP", "キック専用のシンセプラグイン（Windows VST）", "主に psytrance・hitech・darkpsy 向け。プリセット113", "KVR で US$34.73", "△ キックのみ・別ジャンル", "G9"),
]
ws = wb.create_sheet("ジャンル別キット")
ws.cell(row=1, column=1, value="lo-fi と psybient 向けのドラムキットを持つ音源・拡張（" + CHECKED + " 確認）。psybient を明記したドラム音源のキットやプリセットは見つからなかった（NI の Electronica 系エクスパンション一覧にも psychedelic・psy の記載なし [G10]）。").font = Font(name=FONT, italic=True)
header(ws, 2, ["ジャンル", "製品", "開発元", "使う音源（このブックの製品#）", "中身", "価格", "ジャンルの明記", "出典"])
for r, g in enumerate(GENRE, 3):
    body(ws, r, list(g))
for i, w in enumerate([16, 30, 18, 36, 60, 26, 18, 8], 1): ws.column_dimensions[L(i)].width = w
r0 = 3 + len(GENRE) + 1
header(ws, r0, ["番号", "題名", "URL"])
for i, (k, (t, u)) in enumerate(GS.items(), r0 + 1):
    body(ws, i, [k, t, u])
    ws.cell(row=i, column=3).hyperlink = u
    ws.cell(row=i, column=3).font = Font(name=FONT, color="0563C1", underline="single")

used = {s for v in C.values() for s in v[2].split(",")} | {"S24", "S48"}
assert used <= set(S), used - set(S)
print("unused sources:", set(S) - used)
wb.save(OUT)
print("cells", len(C))
